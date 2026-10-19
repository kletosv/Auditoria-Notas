"""Relatório de seis abas; não presume vigência, omissão ou revogação."""
from io import BytesIO
import re
from collections import defaultdict
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

NAMES=['CONSOLIDADO','INCLUSÕES','RETIFICAÇÕES','ONDE CONSTA E PASSE A CONSTAR','TORNAR SEM EFEITO','HISTÓRICO POR MATRÍCULA']

def excel_bytes(docs,hits,differences=None,decisions=None,version='',timeline=None,per_matricula=None,controlled=None,medias=None,cancelamentos=None,base_lookup=None):
    hits=hits if hits is not None else pd.DataFrame()
    timeline=timeline if timeline is not None else pd.DataFrame()
    per_matricula=per_matricula if per_matricula is not None else pd.DataFrame()
    controlled=controlled if controlled is not None else pd.DataFrame()
    medias=medias if medias is not None else pd.DataFrame()
    cancelamentos=cancelamentos if cancelamentos is not None else pd.DataFrame()
    differences=differences if differences is not None else pd.DataFrame()
    sheets={}
    base_lookup=base_lookup or {}
    def base_note(mat,year):
        import re
        key=re.sub(r'\D','',str(mat))
        return str(base_lookup.get(key,{}).get('NOTA '+str(year),'')).strip() if str(year) in ('2019','2020','2021','2022','2023','2024','2025') else ''
    # Consolidado não junta notas alternativas em uma mesma célula: cada evidência tem sua linha.
    records=[]
    if not timeline.empty:
        for _,r in timeline.iterrows():
            records.append({'Matrícula':r.get('Matrícula',''),'Servidor':r.get('Servidor',''),
                'Exercício':r.get('Exercício',''),'Nota Base Geral':base_note(r.get('Matrícula',''),r.get('Exercício','')) or r.get('Nota na Base Geral',''),
                'Nota publicada':r.get('Nota publicada (extração)',''),
                'Código publicado':r.get('Código publicado (extração)',''),
                'Tipo / natureza':r.get('Tipo de ocorrência',''),
                'Alerta de exercício':('EXERCÍCIO INDETERMINADO — NÃO ATRIBUIR NOTA A ANO' if str(r.get('Exercício','')) not in ('2019','2020','2021','2022','2023','2024','2025') else ''),
                'Situação':r.get('Estado documental',''), 'DOE':r.get('DOE',''),
                'Página':r.get('Página',''),'ID evento':r.get('ID do evento','')})
    # Desagregar múltiplos valores extraídos de uma mesma ocorrência.
    expanded=[]
    for rec in records:
        notes=str(rec['Nota publicada']).split(' | ') if rec['Nota publicada'] else ['']
        codes=str(rec['Código publicado']).split(' | ') if rec['Código publicado'] else ['']
        for note in notes:
            for code in codes:
                row=rec.copy();row['Nota publicada']=note;row['Código publicado']=code
                expanded.append(row)
    sheets['CONSOLIDADO']=pd.DataFrame(expanded,columns=['Matrícula','Servidor','Exercício','Nota Base Geral','Nota publicada','Código publicado','Tipo / natureza','Alerta de exercício','Situação','DOE','Página','ID evento'])
    # Inclusão oficial não se confunde com ausência de registro na Base Geral.
    # Sem comprovação da omissão na publicação anterior, não há linha confirmada.
    inclusao_cols=['Matrícula','Servidor','Exercício','O que foi incluído?',
        'Como está na Base Geral','Como deverá ficar na Base Geral',
        'Em qual D.O. consta a inclusão?','A inclusão já consta na Base?',
        'Houve alteração posterior?','Em qual D.O.?','Situação da inclusão',
        'Providência necessária','Cargo','Vínculo','Histórico documental e observações']
    # V19.2: separar candidatos de inclusões com objeto comprovado.
    # Uma palavra-chave ou nota isolada não comprova inclusão nem vigência.
    inclusion_rows=[]
    def _norm_value(v):
        raw=str(v or '').strip().upper()
        if not raw or raw in ('NAN','NONE'): return ''
        if re.fullmatch(r'\d{1,3}(?:[,.]\d{1,2})?',raw):
            return ('%.4f' % float(raw.replace(',','.'))).rstrip('0').rstrip('.')
        return raw
    if not timeline.empty:
        candidates=timeline[timeline['Tipo de ocorrência'].astype(str).str.contains(r'INCLUSAO|INCLUSÃO',case=False,na=False,regex=True)]
        for _,ev in candidates.iterrows():
            mat=str(ev.get('Matrícula','')).strip()
            yr=str(ev.get('Exercício','')).strip()
            trecho=str(ev.get('Trecho da publicação','') or '')
            contexto=str(ev.get('Contexto do ato','') or '')
            full=(contexto+' '+trecho).upper()
            notes=[x.strip() for x in str(ev.get('Nota publicada (extração)','') or '').split(' | ') if x.strip()]
            codes=[x.strip() for x in str(ev.get('Código publicado (extração)','') or '').split(' | ') if x.strip()]
            unique_vals=list(dict.fromkeys(notes+codes))
            year_ok=bool(re.fullmatch(r'20(?:19|2[0-5])',yr))
            # Sem vinculação inequívoca da nota à linha e ao cabeçalho do ato,
            # um valor extraído continua indício, nunca valor a aplicar na Base.
            value_hint=unique_vals[0] if year_ok and len(unique_vals)==1 else ''
            is_media=bool(re.search(r'\bM[ÉE]DIA\b|\bTRI[ÊE]NIO\b',full,re.I))
            has_explicit_object=bool(re.search(r'\bINCLUS[ÃA]O\s+(?:DA\s+)?NOTA\b|\bINCLUIR\s+(?:A\s+)?NOTA\b',full,re.I))
            # V19.3: classifica indícios textuais sem converter classificação em comprovação.
            object_types=[]
            if re.search(r'\bNOTAS?\b|\bPONTUA[CÇ][ÃA]O\b',full,re.I): object_types.append('NOTA/PONTUAÇÃO')
            if re.search(r'\bC[0-9]{3}\b|\bC[ÓO]DIGO\b',full,re.I): object_types.append('CÓDIGO')
            if re.search(r'\bM[ÉE]DIA\b|\bTRI[ÊE]NIO\b|\bQUADRI[ÊE]NIO\b',full,re.I): object_types.append('MÉDIA/PERÍODO')
            if re.search(r'\bMATR[IÍ]CULA\b',full,re.I): object_types.append('MATRÍCULA')
            if re.search(r'\bNOME\b|\bSERVIDOR(?:A)?\b',full,re.I): object_types.append('IDENTIFICAÇÃO DO SERVIDOR')
            object_hint='; '.join(dict.fromkeys(object_types)) if object_types else 'OBJETO NÃO DETECTADO NO TEXTO EXTRAÍDO'
            verified_object=bool(year_ok and value_hint and has_explicit_object and not is_media)
            # A confirmação exige ato/linha/cabeçalho e cobertura posterior, ainda não demonstrados.
            object_label=(f'INDÍCIO: {object_hint}; valor {value_hint}; exercício {yr} — CONFERIR LINHA DO ATO' if verified_object
                          else f'INDÍCIO(S) TEXTUAL(IS): {object_hint} — OBJETO/ALCANCE NÃO COMPROVADOS')
            base=base_note(mat,yr)
            # Comparação é indiciária: não implica que a inclusão foi oficialmente comprovada.
            if verified_object and base:
                comparison=('COINCIDE COM VALOR CANDIDATO — VALIDAR ATO' if _norm_value(base)==_norm_value(value_hint)
                            else 'DIVERGE DO VALOR CANDIDATO — VALIDAR ATO')
            elif not base and year_ok: comparison='BASE SEM VALOR — CONFERIR VÍNCULO E ELEGIBILIDADE'
            else: comparison='PENDENTE — DADOS INSUFICIENTES'
            doe=str(ev.get('DOE','')).strip(); page=str(ev.get('Página','')).strip()
            ref=f'DOE {doe} p. {page}'
            later=str(ev.get('Eventos posteriores possíveis','') or '').strip()
            inclusion_rows.append({'Matrícula':mat,'Servidor':ev.get('Servidor',''),
                'Exercício':yr,'O que foi incluído?':object_label,
                'Como está na Base Geral':base,
                'Como deverá ficar na Base Geral':'NÃO DETERMINADO — ATO E POSTERIORES PENDENTES',
                'Em qual D.O. consta a inclusão?':ref,
                'A inclusão já consta na Base?':comparison,
                'Houve alteração posterior?':('POSSÍVEL — EXAMINAR EVENTOS' if later else 'NÃO COMPROVADO NO LOTE'),
                'Em qual D.O.?':'',
                'Situação da inclusão':'CANDIDATO — NÃO HOMOLOGADO',
                'Providência necessária':'VALIDAR OBJETO, CABEÇALHO, VÍNCULO, ATO ORIGINAL E POSTERIORES',
                'Cargo':'','Vínculo':'',
                'Histórico documental e observações':trecho[:850],
                'Valor extraído (indício)':value_hint,
                'Critério da triagem':('OBJETO SUGERIDO POR TEXTO; LINHA NÃO COMPROVADA' if verified_object else 'SEM OBJETO EXPLÍCITO E UNÍVOCO'),
                'ID do evento':str(ev.get('ID do evento','')),
                'Arquivo de origem':str(ev.get('Arquivo','')),
                'Contexto do ato':contexto[:900],
                'Indícios de objeto no texto':object_hint,
                'Exercício extraído confiável':'SIM — FORMATO/ANO VÁLIDO' if year_ok else 'NÃO — EXERCÍCIO INDETERMINADO',
                'Motivo da pendência':('Ato/linha e publicações posteriores ainda precisam ser confirmados.' if verified_object else 'Não há associação inequívoca entre objeto, matrícula, exercício e valor no texto extraído.')})
    extra_inclusion_cols=['Valor extraído (indício)','Critério da triagem','ID do evento','Arquivo de origem','Contexto do ato','Indícios de objeto no texto','Exercício extraído confiável','Motivo da pendência']
    sheets['INCLUSÕES']=pd.DataFrame(inclusion_rows,columns=inclusao_cols+extra_inclusion_cols)
    ret=[]
    if not controlled.empty:
        for _,r in controlled.iterrows():
            ret.append({'Matrícula':r.get('matricula',''),'Servidor':r.get('nome_depois',''),
                'Exercício':r.get('exercicio',''),'Nota Base Geral':base_note(r.get('matricula',''),r.get('exercicio','')),
                'Nota anterior no DOE':r.get('nota_antes',''),'Nota retificada no DOE':r.get('nota_depois',''),
                'Campo alterado':r.get('campo_alterado',''),'DOE':r.get('doe',''),
                'Edital':r.get('edital',''),'Página':r.get('pagina_impressa',''),
                'Situação':'ALTERAÇÃO IDENTIFICADA; COMPARAR BASE E POSTERIORES'})
    # Fallback conservador: hits podem conter a nota da base, mas não são a Base Geral inteira.
    # A consulta direta ao CSV carregado pelo usuário tem precedência.
    if not hits.empty:
        for r in ret:
            if r['Nota Base Geral']:
                continue
            matches=hits[(hits['Matrícula'].astype(str)==str(r['Matrícula'])) & (hits['Ano'].astype(str)==str(r['Exercício']))]
            vals=list(dict.fromkeys(str(v).strip() for v in matches['Nota na base'] if str(v).strip() and str(v).strip().lower()!='nan'))
            if len(vals)==1:
                r['Nota Base Geral']=vals[0]
    def normalized_note(value):
        raw=str(value).strip()
        if not raw or raw.lower()=='nan': return ''
        if re.fullmatch(r'[0-9]+(?:[.,][0-9]+)?',raw):
            try: return ('%.4f' % float(raw.replace(',','.'))).rstrip('0').rstrip('.')
            except ValueError: pass
        return raw.upper()
    ret_out=[]
    for r in ret:
        base=normalized_note(r['Nota Base Geral']); new=normalized_note(r['Nota retificada no DOE'])
        same=bool(base and new and base==new)
        r['Comparação Base x Retificada']=('COINCIDE COM ESTA RETIFICAÇÃO' if same else
            'DIVERGE — CONFERIR VIGÊNCIA' if base and new else 'NÃO COMPARÁVEL — DADO AUSENTE')
        # Sem cobertura documental posterior comprovada, não declarar SIM ou NÃO.
        r['O que foi alterado?']=r.get('Campo alterado','')
        r['Nota Base']=r.get('Nota Base Geral','')
        r['Nota Retificada']=r.get('Nota retificada no DOE','')
        r['D.O. onde consta a alteração']=r.get('DOE','')
        r['Houve alteração posterior?']='PENDENTE'
        r['Em qual D.O.?']=''
        r['A nota da Base já está corrigida?']='PENDENTE' if base and new else 'PENDENTE — DADOS INSUFICIENTES'
        r['Observação de validação']=('Coincide com a retificação localizada; falta verificar atos posteriores.' if same else
            'Diferença localizada; confirmar cadeia de publicações antes de corrigir a Base Geral.')
        ret_out.append(r)
    ret_cols=['Matrícula','Servidor','Exercício','O que foi alterado?','Nota Base',
        'Nota Retificada','D.O. onde consta a alteração','Houve alteração posterior?',
        'Em qual D.O.?','A nota da Base já está corrigida?',
        'Nota anterior no DOE','Campo alterado','Comparação Base x Retificada',
        'DOE','Edital','Página','Situação','Observação de validação']
    sheets['RETIFICAÇÕES']=pd.DataFrame(ret_out,columns=ret_cols)
    media_rows=[]
    for _,r in medias.iterrows():
        media_rows.append({'Matrícula':r.get('matricula',''),'Servidor':r.get('servidor',''),
            'Triênio':r.get('trienio',''),'Nota '+str(r.get('ano1',''))+' anterior':r.get('antes1',''),
            'Nota '+str(r.get('ano2',''))+' anterior':r.get('antes2',''),
            'Nota '+str(r.get('ano3',''))+' anterior':r.get('antes3',''),
            'Média anterior':r.get('media_antes',''),
            'Nota '+str(r.get('ano1',''))+' corrigida':r.get('depois1',''),
            'Nota '+str(r.get('ano2',''))+' corrigida':r.get('depois2',''),
            'Nota '+str(r.get('ano3',''))+' corrigida':r.get('depois3',''),
            'Média corrigida':r.get('media_depois',''),
            'Campos alterados':r.get('campos_alterados',''),'DOE':r.get('doe',''),
            'Edital':r.get('edital',''),'Publicação de origem':r.get('origem',''),
            'Páginas DOE':r.get('pagina_doe',''),'Situação':r.get('situacao','')})
    sheets['ONDE CONSTA E PASSE A CONSTAR']=pd.DataFrame(media_rows)
    # O objeto atingido deve ser explícito e comprovado pelo texto do ato.
    # Uma mera classificação TORNAR_SEM_EFEITO não comprova NOTA, MÉDIA ou RETIFICAÇÃO.
    t=timeline[timeline['Tipo de ocorrência'].astype(str).str.contains('TORNAR_SEM_EFEITO',na=False)] if not timeline.empty else pd.DataFrame()
    sem_cols=['Matrícula','Servidor','Exercício / Período','O que foi tornado sem efeito?',
        'Nota ou média atingida','D.O. da publicação original','Data e página da publicação original',
        'D.O. que tornou sem efeito','Data e página do cancelamento','Motivo informado no ato',
        'Houve publicação posterior?','Em qual D.O.?','Como está na Base Geral',
        'Como deverá ficar na Base Geral','Providência necessária','Situação da comprovação',
        'Resumo explicativo / Histórico','Ato anterior atingido (comprovado)',
        'Escopo atingido (comprovado)','Arquivo de origem','Página da ocorrência','ID do evento']
    sem_rows=[]
    for _,r in t.iterrows():
        mat=str(r.get('Matrícula','')).strip()
        ano=str(r.get('Exercício','')).strip()
        doe=str(r.get('DOE','')).strip()
        pg=str(r.get('Página','')).strip()
        # Apenas usar campos de comprovação explícitos, nunca deduzir pelo tipo de ocorrência.
        obj=str(r.get('Objeto tornado sem efeito (comprovado)','')).strip()
        if obj.upper() not in ('NOTA','MÉDIA','RETIFICAÇÃO','OUTRO ATO'): obj='PENDENTE DE IDENTIFICAÇÃO'
        ato=str(r.get('Ato anterior atingido (comprovado)','')).strip()
        escopo=str(r.get('Escopo atingido (comprovado)','')).strip()
        valor=str(r.get('Valor atingido (comprovado)','')).strip()
        contexto=(str(r.get('Contexto do ato',''))+' '+str(r.get('Trecho da publicação','')))
        referencias=list(dict.fromkeys(re.findall(r'(?i)\b(?:EDITAL|PORTARIA|RESOLUÇÃO|RESOLUCAO)\s*(?:N[º°.]?\s*)?\d{1,4}\s*/\s*(?:SED/)?20\d{2}',contexto)))
        doe_refs=list(dict.fromkeys(re.findall(r'(?i)\b(?:DOE|D\.?O\.?)\s*(?:N[º°.]?\s*)?1[012][.\s]?\d{3}',contexto)))
        indicios=('Referências textuais a conferir: '+', '.join(referencias[:5]+doe_refs[:5])) if (referencias or doe_refs) else 'Ato anterior não identificado automaticamente'
        comprovado=bool(obj!='PENDENTE DE IDENTIFICAÇÃO' and ato and escopo)
        status='COMPROVADO — CONFERIR POSTERIORES' if comprovado else 'PENDENTE — IDENTIFICAR ATO ANTERIOR E ALCANCE'
        resumo=(f'{obj}: {valor}. Ato anterior: {ato}. Escopo: {escopo}. '
                'Conferir publicações posteriores antes de corrigir a Base Geral.' if comprovado else
                f'O DOE {doe}, página {pg}, contém ocorrência candidata de tornar sem efeito. '
                'Ainda não foi comprovado se atinge nota, média, retificação ou outro ato; não alterar a Base Geral. '+indicios)
        sem_rows.append({'Matrícula':mat,'Servidor':r.get('Servidor',''),
            'Exercício / Período':ano,'O que foi tornado sem efeito?':obj,
            'Nota ou média atingida':valor if comprovado else '',
            'D.O. da publicação original':r.get('DOE original comprovado','') if comprovado else '',
            'Data e página da publicação original':r.get('Data e página original comprovadas','') if comprovado else '',
            'D.O. que tornou sem efeito':doe,
            'Data e página do cancelamento':f'Página {pg}' if pg else '',
            'Motivo informado no ato':r.get('Motivo literal do ato',''),
            'Houve publicação posterior?':'PENDENTE','Em qual D.O.?':'',
            'Como está na Base Geral':base_note(mat,ano),
            'Como deverá ficar na Base Geral':'PENDENTE DE VALIDAÇÃO',
            'Providência necessária':'CONFERIR ATO ANTERIOR E PUBLICAÇÕES POSTERIORES',
            'Situação da comprovação':status,'Resumo explicativo / Histórico':resumo,
            'Ato anterior atingido (comprovado)':ato if comprovado else '',
            'Escopo atingido (comprovado)':escopo if comprovado else '',
            'Arquivo de origem':r.get('Arquivo',''),'Página da ocorrência':pg,
            'ID do evento':r.get('ID do evento','')})
    # Registros com objeto comprovado pelo texto do despacho substituem candidatos
    # genéricos da mesma matrícula/DOE; jamais presumir situação jurídica final.
    if not cancelamentos.empty:
        for _,c in cancelamentos.iterrows():
            mat=str(c['matricula']); doe=str(c['doe']); pg=str(c['pagina_pdf'])
            sem_rows=[r for r in sem_rows if not (str(r['Matrícula'])==mat and str(r['D.O. que tornou sem efeito']).replace('.','')==doe and str(r['Página da ocorrência'])==pg)]
            sem_rows.append({'Matrícula':mat,'Servidor':c['servidor'],
                'Exercício / Período':c['periodo_declarado'] if c['periodo_declarado']==c['cabecalho_periodo'] or not c['cabecalho_periodo'] else 'DIVERGÊNCIA: declarado '+c['periodo_declarado']+' / cabeçalho '+c['cabecalho_periodo'],
                'O que foi tornado sem efeito?':'NOTA — PARTE DE EDITAL',
                'Nota ou média atingida':c['valores'],'D.O. da publicação original':c['doe_anterior'],
                'Data e página da publicação original':'PÁGINA DO ORIGINAL NÃO VERIFICADA',
                'D.O. que tornou sem efeito':doe,'Data e página do cancelamento':'Página PDF '+pg,
                'Motivo informado no ato':'NÃO INFORMADO NO TRECHO',
                'Houve publicação posterior?':'PENDENTE','Em qual D.O.?':'',
                'Como está na Base Geral':base_note(mat,c['periodo_declarado']),
                'Como deverá ficar na Base Geral':'PENDENTE DE VALIDAÇÃO',
                'Providência necessária':'AVALIAR DIVERGÊNCIA DOCUMENTAL' if 'DIVERGÊNCIA' in c['situacao'] else 'CONFERIR ORIGINAL E ATOS POSTERIORES',
                'Situação da comprovação':c['situacao'],
                'Resumo explicativo / Histórico':'Despacho torna sem efeito '+c['escopo'].lower()+' em '+c['ato_anterior']+'. Valores da linha: '+c['valores']+'. Não presume anulação de notas de outros atos.',
                'Ato anterior atingido (comprovado)':c['ato_anterior'],
                'Escopo atingido (comprovado)':c['escopo'],
                'Arquivo de origem':c['arquivo'],'Página da ocorrência':pg,'ID do evento':'CONTROLADO:'+doe+':'+mat})
    sheets['TORNAR SEM EFEITO']=pd.DataFrame(sem_rows,columns=sem_cols)
    hist=timeline.copy()
    if not medias.empty:
        extras=[]
        for _,r in medias.iterrows():
            extras.append({'Matrícula':r['matricula'],'Servidor':r['servidor'],'Exercício':r['trienio'],
                'DOE':r['doe'],'Arquivo':r['arquivo'],'Página':r['pagina_doe'],
                'Tipo de ocorrência':'RETIFICAÇÃO DE MÉDIA DO TRIÊNIO',
                'Natureza do valor':'MÉDIA DE PERÍODO — NÃO TRATAR COMO NOTA ANUAL',
                'Campo afetado':r['campos_alterados'],
                'Onde consta (texto do ato)':f"{r['antes1']}; {r['antes2']}; {r['antes3']}; MÉDIA {r['media_antes']}",
                'Passe a constar (texto do ato)':f"{r['depois1']}; {r['depois2']}; {r['depois3']}; MÉDIA {r['media_depois']}",
                'Estado documental':r['situacao'],'Edital relacionado':r['edital']})
        hist=pd.concat([hist,pd.DataFrame(extras)],ignore_index=True)
    # Histórico executivo: eventos resumidos com referência, nunca notas sem exercício.
    # Matrícula e vínculo só podem ser separados quando vínculo estiver documentado.
    summary=[]
    def clean(v):
        if v is None: return ''
        z=str(v).strip()
        return '' if z.lower() in ('nan','none','nat') else z
    def short_unique(values,limit=800):
        vals=list(dict.fromkeys(clean(v) for v in values if clean(v)))
        out='; '.join(vals)
        return out if len(out)<=limit else out[:limit-3].rsplit('; ',1)[0]+'...'
    def canonical_name(values):
        vals=[clean(x) for x in values if clean(x)]
        if not vals: return ''
        # Mantém uma grafia única, sem concatenar variantes de caixa.
        return max(vals,key=lambda x:(sum(c.islower() for c in x),len(x)))
    def classify(kind):
        x=clean(kind).upper().replace('_',' ')
        if 'TORNAR SEM EFEITO' in x: return 'Tornar sem efeito (alcance não comprovado)'
        if 'MÉDIA' in x or 'MEDIA' in x: return 'Retificação de média' if 'RETIFIC' in x else 'Média de período'
        if 'RETIFIC' in x: return 'Retificação (verificar objeto)'
        if 'INCLUS' in x: return 'Inclusão (verificar omissão anterior)'
        if 'ANULA' in x: return 'Anulação (verificar alcance)'
        if 'REVOGA' in x: return 'Revogação (verificar alcance)'
        return 'Publicação / ocorrência documental'
    if not hist.empty and 'Matrícula' in hist.columns:
        for mat,grp in hist.groupby('Matrícula',dropna=False,sort=False):
            if not clean(mat): continue
            g=grp.copy()
            names=g.get('Servidor',pd.Series('',index=g.index)).tolist()
            periods=[]; event_descriptions=[]; refs=[]; pend=[]
            for _,ev in g.iterrows():
                yr=clean(ev.get('Exercício',''))
                kind=clean(ev.get('Tipo de ocorrência',''))
                doe=clean(ev.get('DOE',''))
                page=clean(ev.get('Página',''))
                label=classify(kind)
                if re.fullmatch(r'20(?:19|2[0-5])',yr): periods.append(yr)
                elif re.fullmatch(r'20\d{2}\s*[–-]\s*20\d{2}',yr): periods.append(yr.replace(' ','').replace('–','-'))
                elif yr and yr.upper()!='INDETERMINADO': pend.append('Conferir período não padronizado')
                else: pend.append('Exercício não identificado em parte das ocorrências')
                ref=f'DOE {doe}'+(f', p. {page}' if page else '') if doe else ''
                if ref: refs.append(ref)
                if label not in event_descriptions: event_descriptions.append(label)
            # Eventos de média recebem contexto explícito; notas anuais só com ano inequívoco.
            media_examples=[]
            if not medias.empty and 'matricula' in medias.columns:
                related=medias[medias['matricula'].astype(str).str.strip()==clean(mat)]
                for _,m in related.iterrows():
                    tri=clean(m.get('trienio',''))
                    old=clean(m.get('media_antes','')); new=clean(m.get('media_depois',''))
                    if tri and old and new:
                        media_examples.append(f'Média {tri}: {old} → {new} (DOE {clean(m.get("doe",""))})')
            period_sorted=sorted(set(periods),key=lambda v:(int(v[:4]),len(v),v))
            obj=short_unique(event_descriptions,240)
            highlights=short_unique(media_examples,300)
            summary_text=(obj+'.' if obj else 'Ocorrências documentais localizadas.')
            if highlights: summary_text+=' '+highlights+'.'
            annual=[]
            if not controlled.empty and 'matricula' in controlled.columns:
                for _,c in controlled[controlled['matricula'].astype(str).str.strip()==clean(mat)].iterrows():
                    annual.append(f"Nota {clean(c.get('exercicio'))}: {clean(c.get('nota_antes'))} → {clean(c.get('nota_depois'))} (DOE {clean(c.get('doe'))})")
            if annual: summary_text+=' '+short_unique(annual,280)+'.'
            if not cancelamentos.empty:
                linked=cancelamentos[cancelamentos['matricula'].astype(str)==clean(mat)]
                for _,c in linked.iterrows():
                    summary_text+=' DOE '+clean(c.get('doe'))+': despacho tornou sem efeito parte de '+clean(c.get('ato_anterior'))+' ('+clean(c.get('valores'))+').'
                    if 'DIVERGÊNCIA' in clean(c.get('situacao')):
                        summary_text+=' Divergência entre período declarado e cabeçalho; pendente de avaliação documental.'  
            summary_text+=' Vigência e alterações posteriores ainda não confirmadas no universo documental.'
            if any('Tornar sem efeito' in x for x in event_descriptions):
                pend.append('Identificar ato e valor tornados sem efeito')
            pend.append('Verificar cobertura de DOE principal e suplementos e atos posteriores')
            summary.append({
                'Matrícula':clean(mat),
                'Servidor':canonical_name(names),
                'Exercícios com ocorrências':short_unique(period_sorted,180),
                'O que foi encontrado?':obj,
                'Resumo do histórico':summary_text[:650],
                'Situação da Base Geral':'PENDENTE DE CONFERÊNCIA',
                'Providência necessária':'CONFERIR ATOS E COMPARAR COM A BASE',
                'Diários e páginas':short_unique(refs,1500),
                'Pendências':short_unique(pend,400),
            })
    sheets['HISTÓRICO POR MATRÍCULA']=pd.DataFrame(summary,columns=[
        'Matrícula','Servidor','Exercícios com ocorrências','O que foi encontrado?',
        'Resumo do histórico','Situação da Base Geral','Providência necessária',
        'Diários e páginas','Pendências'])

    buf=BytesIO()
    with pd.ExcelWriter(buf,engine='openpyxl') as writer:
        for name in NAMES:
            frame=sheets[name].copy().fillna('')
            if frame.empty and len(frame.columns)==0:frame=pd.DataFrame(columns=['Sem registros comprovados'])
            for col in frame.columns:
                if 'Matrícula' in str(col) or str(col)=='DOE':frame[col]=frame[col].astype(str)
            frame.to_excel(writer,sheet_name=name,index=False)
            ws=writer.sheets[name];ws.freeze_panes='A2';ws.sheet_view.showGridLines=False
            for c in ws[1]:
                c.fill=PatternFill('solid',fgColor='17324D');c.font=Font(bold=True,color='FFFFFF');c.alignment=Alignment(wrap_text=True)
            ws.row_dimensions[1].height=32
            for i,col in enumerate(frame.columns,1):
                ws.column_dimensions[get_column_letter(i)].width=min(43,max(17,len(str(col))+3))
                if 'Matrícula' in str(col) or str(col)=='DOE':
                    for row in ws.iter_rows(min_row=2,min_col=i,max_col=i):row[0].number_format='@'
            if name=='HISTÓRICO POR MATRÍCULA':
                for letter,width in {'A':17,'B':35,'C':25,'D':48,'E':90,'F':29,'G':42,'H':65,'I':65}.items():
                    ws.column_dimensions[letter].width=width
                for row in ws.iter_rows(min_row=2):
                    for cell in row:
                        cell.alignment=Alignment(wrap_text=True,vertical='top')
                    ws.row_dimensions[row[0].row].height=65
                ws.freeze_panes='C2'
            if name=='TORNAR SEM EFEITO':
                obj_col=4
                ws.column_dimensions['D'].width=34
                ws.column_dimensions['P'].width=49
                ws.column_dimensions['Q'].width=76
                for row in ws.iter_rows(min_row=2):
                    cell=row[obj_col-1]
                    is_pending='PENDENTE' in str(cell.value).upper()
                    cell.fill=PatternFill('solid',fgColor='FFF2CC' if is_pending else 'FCE4D6')
                    cell.font=Font(bold=True,color='9C5700' if is_pending else '9C0006')
                    cell.alignment=Alignment(wrap_text=True,vertical='center')
                    row[16].alignment=Alignment(wrap_text=True,vertical='center')
                    ws.row_dimensions[cell.row].height=42
            if ws.max_row>1:ws.auto_filter.ref=ws.dimensions
    return buf.getvalue()
